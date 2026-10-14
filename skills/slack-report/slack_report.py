"""확인 결과를 Slack 채널에 보낸다: 요약 글 + 사용한 쿼리 + 그래프(PNG) + 표(엑셀).

  쿼리 검사(실행 전에): run.sh --check-sql query.sql
  만들기·보내기:       run.sh --title "센터별 회원 수 (2026-10-13)" --data rows.csv --sql query.sql \
                          [--db "DB 이름"] [--summary "요약 글"] [--chart bar|line --x 열 --y 열[,열] [--top 10]] \
                          (--dry-run | --send | --channel C0123ABCD)

--data 는 머리줄이 있는 CSV(UTF-8), --sql 은 그 결과를 낸 쿼리 그대로다. 같은 폴더에 <제목>.png·<제목>.xlsx·<제목>.sql 을 만든다.
쿼리가 조회(SELECT 등)가 아니면 아무것도 만들거나 보내지 않고 멈춘다.
--top N 은 그래프에만 세로축 첫 열이 큰 순으로 N개를 그린다(엑셀에는 전부 들어간다).
쿼리는 글에 넣지 않고 <제목>.sql 로 붙인다(Slack에서 접혀 보인다).
--dry-run 이면 파일만 만들고 Slack에는 보내지 않는다(보내기 전에 확인용).
설정은 SLACK_BOT_TOKEN·SLACK_CHANNEL 환경 변수, 없으면 ~/.config/slack-report/.env 에서 읽는다(.env.template 을 복사해 채운다).
--send 는 SLACK_CHANNEL 로 보낸다.
봇 권한은 chat:write, files:write 이고, 봇을 채널에 초대해 둬야 한다.
"""
import argparse
import csv
import json
import os
import re
import sys
import urllib.parse
import urllib.request
from pathlib import Path

# 플러그인 폴더는 업데이트 때 새로 깔리므로 설정은 폴더 밖에 둔다.
ENV_FILE = Path.home() / '.config' / 'slack-report' / '.env'
# 한글 글꼴. 있는 것 가운데 앞의 것을 쓴다(macOS, Windows, 리눅스 순).
FONTS = ['AppleGothic', 'Apple SD Gothic Neo', 'Malgun Gothic', 'NanumGothic', 'Noto Sans CJK KR', 'Noto Sans KR']
# 엑셀에 숫자로 넣을 꼴. 앞자리 0(번호)·16자리 이상(엑셀이 끝자리를 잃는다)·전화번호는 글자 그대로 둔다.
NUMBER = re.compile(r'-?(0|[1-9]\d{0,14})(\.\d+)?')
# 조회 쿼리 검사. 실수를 거르는 장치이고, 실행을 막는 진짜 보장은 메타베이스가 읽기 전용 DB 계정으로 붙는 것이다.
READ_START = re.compile(r'(SELECT|WITH|SHOW|DESCRIBE|DESC|EXPLAIN)\b', re.I)
NOT_READ = re.compile(r'\b(INSERT|UPDATE|DELETE|MERGE|UPSERT|CREATE|ALTER|DROP|TRUNCATE|RENAME|GRANT|REVOKE|CALL|EXEC|EXECUTE'
                      r'|LOCK|SHARE|INTO|LOAD_FILE|GET_LOCK|NEXTVAL|SETVAL|PG_TERMINATE_BACKEND|PG_CANCEL_BACKEND|DBLINK_EXEC)\b', re.I)
FOOTER = 'made by hudson and metabase mcp'  # 모든 글 끝에 작은 글씨(context 블록)로 붙인다
QUOTED = re.compile(r"'(?:[^'\\]|\\.|'')*'|\"(?:[^\"\\]|\\.)*\"|`[^`]*`")


def query_problem(sql):
    """조회 쿼리가 아니면 그 까닭을, 맞으면 None 을 돌려준다. 따옴표 안 글자('UPDATE' 같은 값)는 보지 않는다."""
    body = sql.strip().rstrip(';')
    if ';' in body:
        return '문장이 여러 개입니다(; 는 맨 끝에만 쓸 수 있습니다).'
    code = QUOTED.sub("''", body)
    if not READ_START.match(code.lstrip('( \t\r\n')):
        return '조회(SELECT·WITH·SHOW·DESCRIBE·EXPLAIN)로 시작하지 않습니다. 앞에 주석도 쓰지 않습니다.'
    if m := NOT_READ.search(code):
        return f'조회가 아닌 말({m.group(1).upper()})이 들어 있습니다.'
    return None


def number(text):
    t = str(text).replace(',', '').strip()
    if not NUMBER.fullmatch(t):
        return None
    return float(t) if '.' in t else int(t)


def read_table(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        rows = list(csv.reader(f))
    if not rows:
        sys.exit(f'{path} 가 비어 있습니다.')
    return rows[0], rows[1:]


def write_xlsx(header, rows, out):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = '결과'
    ws.append(header)
    for row in rows:
        ws.append([n if (n := number(v)) is not None else v for v in row])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, column in enumerate(zip(header, *rows), 1):
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(len(str(v)) for v in column) * 1.8 + 2)
    ws.freeze_panes = 'A2'
    wb.save(out)


def write_chart(header, rows, kind, x, ys, title, out, top=None):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter
    for name in [x, *ys]:
        if name not in header:
            sys.exit(f'"{name}" 열이 없습니다. 있는 열: {", ".join(header)}')
    if top:
        first = header.index(ys[0])
        rows = sorted(rows, key=lambda r: number(r[first]) or 0, reverse=True)[:top]
        title += f' — 상위 {len(rows)}'
    from matplotlib import font_manager
    installed = {f.name for f in font_manager.fontManager.ttflist}
    plt.rcParams['font.family'] = next((f for f in FONTS if f in installed), 'sans-serif')
    plt.rcParams['axes.unicode_minus'] = False
    n = len(rows)
    fig, ax = plt.subplots(figsize=(10, 5.5), dpi=150)
    width = 0.8 / len(ys)
    for k, y in enumerate(ys):
        values = [number(row[header.index(y)]) or 0 for row in rows]
        if kind == 'bar':
            bars = ax.bar([i + (k - (len(ys) - 1) / 2) * width for i in range(n)], values, width, label=y)
            if n * len(ys) <= 30:
                ax.bar_label(bars, labels=[f'{v:,.0f}' if float(v).is_integer() else f'{v:,.2f}' for v in values], fontsize=8)
        else:
            ax.plot(range(n), values, marker='o' if n <= 40 else None, label=y)
    # 가로축 이름은 20개 안팎만 보인다(날짜가 많으면 겹친다).
    ticks = list(range(0, n, max(1, n // 20)))
    ax.set_xticks(ticks, [rows[i][header.index(x)] for i in ticks], rotation=45 if len(ticks) > 6 else 0, ha='right' if len(ticks) > 6 else 'center')
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f'{v:,.0f}' if abs(v) >= 1000 else f'{v:g}'))
    ax.set_title(title)
    ax.grid(axis='y', alpha=0.3)
    if len(ys) > 1:
        ax.legend()
    fig.tight_layout()
    fig.savefig(out)
    plt.close(fig)


def escape(text):
    """Slack 글에서 &, <, > 는 링크·멘션 기호라 바꿔 써야 한다(쿼리의 < 비교가 깨지지 않게)."""
    return text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def message(title, summary, sql, db):
    text = f'*{escape(title)}*' + (f'\n{escape(summary)}' if summary else '')
    if sql:
        # Slack 글에는 접기가 없어서, 쿼리는 .sql 파일로 붙인다(미리보기 몇 줄만 보이고 누르면 펼쳐진다).
        text += '\n\n*사용한 쿼리*' + (f' (DB: {escape(db)})' if db else '') + ' — 붙인 .sql 파일을 누르면 펼쳐집니다.'
    return text


def slack(method, token, **fields):
    req = urllib.request.Request('https://slack.com/api/' + method, urllib.parse.urlencode(fields).encode(), {'Authorization': f'Bearer {token}'})
    with urllib.request.urlopen(req, timeout=30) as res:
        reply = json.load(res)
    if not reply.get('ok'):
        sys.exit(f'Slack {method} 실패: {reply.get("error")}')
    return reply


def upload(token, path):
    """파일을 Slack에 올리고 파일 번호를 돌려준다. 채널에는 send()의 마지막 호출에서 함께 나타난다."""
    reply = slack('files.getUploadURLExternal', token, filename=path.name, length=path.stat().st_size)
    req = urllib.request.Request(reply['upload_url'], path.read_bytes(), {'Content-Type': 'application/octet-stream'})
    urllib.request.urlopen(req, timeout=60).close()
    return reply['file_id']


def send(token, channel, text, files):
    # 작은 글씨는 context 블록으로만 된다. 파일과 함께 보낼 때 initial_comment 를 주면 blocks 가 무시되므로 blocks 만 준다.
    blocks = json.dumps([{'type': 'section', 'text': {'type': 'mrkdwn', 'text': text}},
                         {'type': 'context', 'elements': [{'type': 'mrkdwn', 'text': FOOTER}]}], ensure_ascii=False)
    if not files:
        slack('chat.postMessage', token, channel=channel, text=text, blocks=blocks)
        return
    uploaded = [{'id': upload(token, f), 'title': f.name} for f in files]
    slack('files.completeUploadExternal', token, files=json.dumps(uploaded, ensure_ascii=False), channel_id=channel, blocks=blocks)


def main():
    p = argparse.ArgumentParser(description='확인 결과를 Slack 채널에 보낸다.')
    p.add_argument('--check-sql', type=Path, help='조회 쿼리인지 검사만 한다(실행 전에)')
    p.add_argument('--title')
    p.add_argument('--summary', default='')
    p.add_argument('--data', type=Path, help='머리줄이 있는 CSV')
    p.add_argument('--sql', type=Path, help='--data 를 낸 쿼리 파일')
    p.add_argument('--db', default='', help='쿼리를 실행한 DB 이름')
    p.add_argument('--chart', choices=['bar', 'line'])
    p.add_argument('--x', help='가로축 열 이름')
    p.add_argument('--y', help='세로축 열 이름(쉼표로 여러 개)')
    p.add_argument('--top', type=int, help='그래프에만 큰 순으로 N개')
    p.add_argument('--send', action='store_true', help='설정 파일의 기본 채널로 보낸다')
    p.add_argument('--channel', help='다른 채널로 보낼 때 채널 ID(C로 시작)')
    p.add_argument('--dry-run', action='store_true', help='파일만 만들고 보내지 않는다')
    a = p.parse_args()
    if a.check_sql:
        problem = query_problem(a.check_sql.read_text(encoding='utf-8'))
        if problem:
            sys.exit(f'실행하면 안 되는 쿼리입니다: {problem}')
        print('조회 쿼리입니다.')
        return
    if not a.title:
        p.error('--title 이 없습니다.')
    if a.data and not a.sql:
        p.error('--data 에는 그 결과를 낸 쿼리(--sql)가 함께 있어야 합니다.')
    if a.chart and not (a.data and a.x and a.y):
        p.error('--chart 에는 --data, --x, --y 가 함께 있어야 합니다.')
    if not (a.dry_run or a.send or a.channel):
        p.error('--send(기본 채널) 나 --channel 이 없습니다(확인만 하려면 --dry-run).')
    sql = a.sql.read_text(encoding='utf-8') if a.sql else ''
    if sql and (problem := query_problem(sql)):
        sys.exit(f'보내지 않습니다. 조회 쿼리가 아닙니다: {problem}')
    name = re.sub(r'[\\/:*?"<>|]', '_', a.title)
    files = []
    if a.data:
        header, rows = read_table(a.data)
        if a.chart:
            files.append(a.data.parent / f'{name}.png')
            write_chart(header, rows, a.chart, a.x, a.y.split(','), a.title, files[-1], a.top)
        files.append(a.data.parent / f'{name}.xlsx')
        write_xlsx(header, rows, files[-1])
    if sql:
        files.append(a.sql.parent / f'{name}.sql')
        files[-1].write_text(sql, encoding='utf-8')
    text = message(a.title, a.summary, sql, a.db)
    if a.dry_run:
        print('만든 파일:', *files, sep='\n  ')
        print('보낼 글:\n' + text + f'\n(작은 글씨) {FOOTER}')
        return
    env = {}
    if ENV_FILE.exists():
        env = {k.strip(): v.strip() for k, v in (line.split('=', 1) for line in ENV_FILE.read_text(encoding='utf-8').splitlines()
                                                  if '=' in line and not line.lstrip().startswith('#'))}
    env.update({k: os.environ[k] for k in ('SLACK_BOT_TOKEN', 'SLACK_CHANNEL') if os.environ.get(k)})
    token, channel = env.get('SLACK_BOT_TOKEN'), a.channel or env.get('SLACK_CHANNEL')
    if not token or not channel:
        template = Path(__file__).parent / '.env.template'
        sys.exit(f'토큰이나 채널이 없습니다. 아래처럼 설정 파일을 만들고 값을 채우세요.\n'
                 f'  mkdir -p {ENV_FILE.parent} && cp "{template}" {ENV_FILE} && chmod 600 {ENV_FILE}')
    send(token, channel, text, files)
    print(f'보냈습니다: 채널 {channel}, 파일 {len(files)}개')


if __name__ == '__main__':
    main()
