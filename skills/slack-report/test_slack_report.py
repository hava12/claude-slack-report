"""slack_report.py 확인(Slack에는 보내지 않는다): ~/.cache/slack-report/venv/bin/python test_slack_report.py (run.sh 를 한 번 실행한 뒤)"""
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from slack_report import number, query_problem

HERE = Path(__file__).parent

assert number('1,234') == 1234 and number('12.5') == 12.5 and number('-3') == -3
assert number('007') is None and number('010-1234-5678') is None and number('1234567890123456') is None and number('') is None

ok = ['SELECT * FROM center LIMIT 5', 'with a as (select 1) select * from a;', '(SELECT 1) UNION (SELECT 2)',
      "SELECT REPLACE(name, ' ', '') FROM center WHERE memo LIKE '%delete%' AND type = 'UPDATE'",
      'SELECT updated_at, is_delete, created_at FROM center', 'SHOW GRANTS', 'EXPLAIN SELECT 1']
bad = ["UPDATE center SET name = 'x'", 'DELETE FROM center', 'SELECT 1; DROP TABLE center',
       'WITH d AS (DELETE FROM center RETURNING *) SELECT * FROM d', 'SELECT * FROM center FOR UPDATE',
       "SELECT * INTO OUTFILE '/tmp/x' FROM center", '-- 설명\nSELECT 1', 'SELECT 1 /*! ; DELETE FROM t */',
       "SELECT GET_LOCK('a', 10)"]
assert not [q for q in ok if query_problem(q)], [q for q in ok if query_problem(q)]
assert not [q for q in bad if not query_problem(q)], [q for q in bad if not query_problem(q)]

with tempfile.TemporaryDirectory() as d:
    data, sql = Path(d) / 'rows.csv', Path(d) / 'q.sql'
    data.write_text('센터,회원 수,번호\n강남점,1200,007\n판교점,950,010\n분당점,300,011\n', encoding='utf-8')
    sql.write_text("SELECT name, cnt\nFROM t\nWHERE cnt > 100 AND type = 'UPDATE'\n", encoding='utf-8')
    cmd = [sys.executable, HERE / 'slack_report.py', '--title', '센터별 회원 수', '--data', data, '--sql', sql,
           '--chart', 'bar', '--x', '센터', '--y', '회원 수', '--top', '2', '--dry-run']
    run = subprocess.run(cmd, capture_output=True, text=True)
    assert run.returncode == 0, run.stderr
    assert 'made by hudson and metabase mcp' in run.stdout, run.stdout  # 끝의 작은 글씨
    assert 'WHERE cnt' not in run.stdout, run.stdout  # 쿼리는 글이 아니라 파일로 붙는다(Slack에서 접혀 보인다)
    assert (Path(d) / '센터별 회원 수.sql').read_text(encoding='utf-8') == sql.read_text(encoding='utf-8')
    ws = load_workbook(Path(d) / '센터별 회원 수.xlsx').active
    assert ws.max_row == 4 and [c.value for c in ws[2]] == ['강남점', 1200, '007']  # 그래프는 상위 2개, 엑셀은 전부
    assert (Path(d) / '센터별 회원 수.png').stat().st_size > 10000
    sql.write_text('DELETE FROM t', encoding='utf-8')
    run = subprocess.run(cmd, capture_output=True, text=True)
    assert run.returncode != 0 and '조회 쿼리가 아닙니다' in run.stderr, run.stderr
print('통과')
